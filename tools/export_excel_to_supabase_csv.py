import csv
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "SHK - DATA KHAIRAT KEMATIAN.xlsx"
OUT_DIR = ROOT / "supabase-import"


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def money(value):
    value = clean(value)
    if not value:
        return ""
    if value in {"-", "TIADA BAYARAN", "TIADA INFO"}:
        return ""
    try:
        return str(float(value.replace(",", "")))
    except ValueError:
        return ""


def integer(value):
    value = clean(value)
    if not value or value in {"-", "TIADA BAYARAN", "TIADA INFO"}:
        return ""
    try:
        return str(int(float(value)))
    except ValueError:
        return ""


def write_csv(name, headers, rows):
    OUT_DIR.mkdir(exist_ok=True)
    with (OUT_DIR / name).open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def main():
    workbook = openpyxl.load_workbook(WORKBOOK, data_only=True)

    keahlian = workbook["KEAHLIAN"]
    members = []
    payments = []

    payment_map = [
        (2019, 8, None),
        (2020, 10, 9),
        (2021, 12, 11),
        (2022, 14, 13),
        (2023, 16, 15),
        (2024, 18, 17),
        (2025, 20, 19),
        (2026, 22, 21),
        (2027, 24, 23),
        (2028, 25, None),
        (2029, 26, None),
        (2030, 27, None),
    ]

    for row in keahlian.iter_rows(min_row=5, values_only=True):
        member_no = clean(row[1])
        name = clean(row[2])
        if not member_no and not name:
            continue

        members.append({
            "member_no": member_no,
            "member_name": name,
            "ic_no": clean(row[3]),
            "phone": clean(row[4]),
            "address": clean(row[5]),
            "dependant_count": integer(row[6]),
            "registration_year": clean(row[27]),
            "registration_member_amount": money(row[28]),
            "registration_dependant_amount": money(row[29]),
            "arrears_amount": money(row[30]),
            "source_sheet": "KEAHLIAN",
        })

        for year, amount_col, receipt_col in payment_map:
            amount = money(row[amount_col - 1])
            receipt_no = clean(row[receipt_col - 1]) if receipt_col else ""
            if amount or receipt_no:
                payments.append({
                    "member_no": member_no,
                    "member_name": name,
                    "payment_year": year,
                    "amount": amount,
                    "receipt_no": receipt_no,
                    "source_sheet": "KEAHLIAN",
                })

    tanggungan = workbook["TANGGUNGAN 1"]
    dependants = []
    dependant_columns = [
        (1, 7, 8),
        (2, 9, 10),
        (3, 11, 12),
        (4, 13, 14),
        (5, 16, 17),
        (6, 18, 19),
        (7, 20, 21),
        (8, 22, 23),
        (9, 24, 25),
        (10, 26, 27),
    ]

    for row in tanggungan.iter_rows(min_row=5, values_only=True):
        member_no = clean(row[1])
        member_name = clean(row[2])
        if not member_no and not member_name:
            continue

        for position, name_col, relation_col in dependant_columns:
            dependant_name = clean(row[name_col - 1])
            relationship = clean(row[relation_col - 1])
            if dependant_name or relationship:
                dependants.append({
                    "member_no": member_no,
                    "member_name": member_name,
                    "position_no": position,
                    "dependant_name": dependant_name,
                    "dependant_ic": "",
                    "relationship": relationship,
                    "source_sheet": "TANGGUNGAN 1",
                })

    pindah = workbook["PINDAHBERHENTI"]
    inactive = []
    for row in pindah.iter_rows(min_row=4, values_only=True):
        member_no = clean(row[1])
        name = clean(row[2])
        if not member_no and not name:
            continue
        inactive.append({
            "member_no": member_no,
            "member_name": name,
            "phone": clean(row[4]),
            "address": clean(row[5]),
            "details": clean(row[6]),
            "source_sheet": "PINDAHBERHENTI",
        })

    write_csv("members.csv", list(members[0].keys()), members)
    write_csv("member_yearly_payments.csv", list(payments[0].keys()), payments)
    write_csv("member_dependants.csv", list(dependants[0].keys()), dependants)
    write_csv("inactive_members.csv", list(inactive[0].keys()), inactive)

    print(f"Exported {len(members)} members")
    print(f"Exported {len(payments)} yearly payment rows")
    print(f"Exported {len(dependants)} dependant rows")
    print(f"Exported {len(inactive)} inactive rows")


if __name__ == "__main__":
    main()
